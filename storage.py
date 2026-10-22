"""
Storage. Every function writes immediately, so the app autosaves on each change.

Where the data lives:
  • DATABASE_URL set (Streamlit secret or environment variable) → that Postgres database (e.g. Neon).
    Data survives restarts and redeploys — use this for the public app.
  • Otherwise → data/throwing.db, a SQLite file next to the app (fine on your own computer).
"""
from __future__ import annotations

import os
import re
import threading
import time
import unicodedata
import uuid
from datetime import date, datetime

import pandas as pd
from sqlalchemy import create_engine, text

HERE = os.path.dirname(os.path.abspath(__file__))
SQLITE_PATH = os.environ.get("THROW_DB", os.path.join(HERE, "data", "throwing.db"))

LOG_FIELDS = ["date", "session_type", "distance", "start_time", "end_time", "actual_throws",
              "velo_low", "velo_high", "warm_throws", "warm_velo_low", "warm_velo_high",
              "sixty_throws", "sixty_velo_low", "sixty_velo_high", "notes"]
NUM_FIELDS = {"distance", "actual_throws", "velo_low", "velo_high", "warm_throws", "warm_velo_low",
              "warm_velo_high", "sixty_throws", "sixty_velo_low", "sixty_velo_high"}
BU_FIELDS = ["start", "first_pen", "pen_pitches", "pen_velo_low", "pen_velo_high", "starting_point", "max_distance", "volume"]

# 2026 pitching staff (jersey, name, position, year, bats/throws) — only used if no backup is found on first run
DEFAULT_ROSTER = [
    ("30", "Max Bettis", "RHP", "So.", "R/R"), ("33", "Sam Bettis", "RHP", "So.", "R/R"),
    ("26", "Noah Gartner", "RHP", "So.", "R/R"), ("31", "Cooper Jesperson", "RHP", "So.", "R/R"),
    ("23", "Jonah Reich", "RHP", "5th", "R/R"), ("32", "Jackson Baker", "RHP", "Sr.", "R/R"),
    ("7", "Trenton Buckley", "RHP", "Jr.", "R/R"), ("17", "Stevie Doty", "RHP", "Jr.", "R/R"),
    ("99", "Colin Driffill", "RHP", "Fr.", "R/R"), ("37", "Jackson Hulcher", "RHP", "Sr.", "R/R"),
    ("35", "Griffin Lewis", "RHP", "Jr.", "R/R"), ("50", "Max Nantais-Vlahovich", "RHP", "Fr.", "R/R"),
    ("28", "Joseph Pereira", "RHP", "Fr.", "L/R"), ("18", "Stocton Timbrook", "RHP", "Jr.", "R/R"),
    ("15", "Grady Westphal", "RHP", "So.", "R/R"), ("4", "Cordell Clinkingbeard", "RHP/INF", "Fr.", "R/R"),
    ("16", "Kyler Horsman", "RHP/OF", "Sr.", "R/R"), ("9", "Donte Lewis", "UTL", "Jr.", "R/R"),
    ("41", "Adam Arther", "LHP", "Sr.", "L/L"), ("0", "Kaden Dean", "LHP", "Fr.", "L/L"),
    ("25", "Grant Meert", "LHP", "Sr.", "L/L"),
]
# The workbook's Build-Up tabs and Roster use last names for these three
NAME_ALIASES = {"lewis": "griffin-lewis", "doty": "stevie-doty", "dean": "kaden-dean"}


def slugify(name: str) -> str:
    s = unicodedata.normalize("NFKD", str(name)).lower()
    s = re.sub(r"[^\w\s-]", "", s).strip()
    return re.sub(r"-+", "-", re.sub(r"[\s_]+", "-", s))


def database_url() -> str | None:
    url = os.environ.get("DATABASE_URL")
    if not url:
        try:
            import streamlit as st
            url = st.secrets.get("DATABASE_URL")
        except Exception:
            url = None
    if url:
        url = url.strip()
        for prefix in ("postgres://", "postgresql://"):
            if url.startswith(prefix):     # use the psycopg2 driver (requirements.txt: psycopg2-binary)
                url = "postgresql+psycopg2://" + url[len(prefix):]
    return url or None


def where_saved() -> str:
    url = database_url()
    if url:
        host = re.sub(r"^.*@", "", url).split("/")[0].split(":")[0]
        return f"Postgres database ({host})"
    return "a file on this computer (data/throwing.db)"


_ENGINE = None
_LOCK = threading.Lock()
_CACHE: dict = {}


def engine():
    global _ENGINE
    if _ENGINE is None:
        with _LOCK:
            if _ENGINE is None:
                url = database_url()
                if url:
                    eng = create_engine(url, pool_pre_ping=True, pool_recycle=280)
                else:
                    os.makedirs(os.path.dirname(SQLITE_PATH), exist_ok=True)
                    eng = create_engine(f"sqlite:///{SQLITE_PATH}")
                _create_tables(eng)
                _ENGINE = eng
    return _ENGINE


def _create_tables(eng):
    cols = ", ".join(f + (" DOUBLE PRECISION" if f in NUM_FIELDS else " TEXT") for f in LOG_FIELDS)
    with eng.begin() as c:
        c.execute(text("""CREATE TABLE IF NOT EXISTS players(
            id TEXT PRIMARY KEY, name TEXT, number TEXT, position TEXT, year TEXT, bt TEXT,
            avg_top_velo DOUBLE PRECISION, peak_velo DOUBLE PRECISION, active INTEGER DEFAULT 1, sort INTEGER,
            target_type TEXT DEFAULT 'Game', target_date TEXT, planned_pitches DOUBLE PRECISION, planned_innings DOUBLE PRECISION)"""))
        c.execute(text(f"CREATE TABLE IF NOT EXISTS throws(id TEXT PRIMARY KEY, player TEXT, created DOUBLE PRECISION, "
                       f"updated DOUBLE PRECISION, {cols})"))
        c.execute(text("CREATE TABLE IF NOT EXISTS overrides(player TEXT, day TEXT, day_type TEXT, PRIMARY KEY(player, day))"))
        c.execute(text("""CREATE TABLE IF NOT EXISTS buildup(player TEXT PRIMARY KEY, start TEXT, first_pen TEXT,
            pen_pitches DOUBLE PRECISION, pen_velo_low DOUBLE PRECISION, pen_velo_high DOUBLE PRECISION,
            starting_point TEXT, max_distance DOUBLE PRECISION, volume DOUBLE PRECISION)"""))
        c.execute(text("""CREATE TABLE IF NOT EXISTS buildup_days(player TEXT, day TEXT, off INTEGER,
            throws DOUBLE PRECISION, distance DOUBLE PRECISION, velo_low DOUBLE PRECISION, velo_high DOUBLE PRECISION,
            PRIMARY KEY(player, day))"""))
        c.execute(text("CREATE INDEX IF NOT EXISTS throws_player ON throws(player)"))


def _q(sql, params=None):
    with engine().connect() as c:
        return c.execute(text(sql), params or {}).fetchall()


def _x(sql, params=None):
    """Write, then drop the read cache (one server process serves every user, so this stays consistent)."""
    with engine().begin() as c:
        if isinstance(params, list):
            c.execute(text(sql), params)
        else:
            c.execute(text(sql), params or {})
    _CACHE.clear()


def _df(sql, params=None):
    with engine().connect() as c:
        return pd.read_sql(text(sql), c, params=params or {})


def _cached(key, fn):
    if key not in _CACHE:
        _CACHE[key] = fn()
    v = _CACHE[key]
    return v.copy() if hasattr(v, "copy") else v


def _clean(field, v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if field in ("date", "day", "start", "first_pen", "target_date"):
        if isinstance(v, datetime):
            v = v.date()
        return v.isoformat() if isinstance(v, date) else (str(v)[:10] or None)
    if field in ("start_time", "end_time"):
        return v.strftime("%H:%M") if hasattr(v, "strftime") else (str(v).strip()[:5] or None)
    if field in NUM_FIELDS or field in ("pen_pitches", "pen_velo_low", "pen_velo_high", "max_distance", "volume",
                                         "throws", "distance", "velo_low", "velo_high", "avg_top_velo", "peak_velo",
                                         "planned_pitches", "planned_innings"):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None
    s = str(v).strip()
    return s or None


# ───────── players ─────────
def load_players(include_inactive=False) -> pd.DataFrame:
    def go():
        return _df("SELECT * FROM players" + ("" if include_inactive else " WHERE active=1") + " ORDER BY sort, name")
    df = _cached(("players", include_inactive), go)
    df["target_date"] = [date.fromisoformat(v) if isinstance(v, str) and v else None for v in df["target_date"]]
    df["target_date"] = df["target_date"].astype(object)
    return df


def get_player(pid) -> dict | None:
    df = load_players(include_inactive=True)
    hit = df[df["id"] == pid]
    if hit.empty:
        return None
    r = hit.iloc[0].to_dict()
    for k in ("avg_top_velo", "peak_velo", "planned_pitches", "planned_innings"):
        if pd.isna(r.get(k)):
            r[k] = None
    if not isinstance(r.get("target_date"), date) or pd.isna(r.get("target_date")):
        r["target_date"] = None
    return r


def resolve_player(name) -> str | None:
    """Full name, id, or (for the workbook's short names) a last name → player id."""
    if not name:
        return None
    pid = slugify(name)
    df = load_players(include_inactive=True)
    if pid in set(df["id"]):
        return pid
    if pid in NAME_ALIASES and NAME_ALIASES[pid] in set(df["id"]):
        return NAME_ALIASES[pid]
    hits = [i for i, n in zip(df["id"], df["name"]) if str(n).split()[-1].lower() == str(name).strip().lower()]
    return hits[0] if len(hits) == 1 else None


def add_player(name, number="", position="RHP", year=None, bt=None, avg_top=None, peak=None) -> str:
    pid = slugify(name)
    existing = _q("SELECT 1 FROM players WHERE id=:id", {"id": pid})
    if existing:  # re-adding a removed player brings his log back
        _x("UPDATE players SET active=1, name=:n, number=:num, position=:pos WHERE id=:id",
           {"n": name.strip(), "num": number, "pos": position, "id": pid})
        if avg_top is not None:
            update_player(pid, avg_top_velo=avg_top)
    else:
        n = _q("SELECT COALESCE(MAX(sort), 0) + 1 FROM players")[0][0]
        _x("""INSERT INTO players(id, name, number, position, year, bt, avg_top_velo, peak_velo, active, sort, target_type)
              VALUES (:id,:n,:num,:pos,:yr,:bt,:avg,:peak,1,:s,'Game')""",
           {"id": pid, "n": name.strip(), "num": number, "pos": position, "yr": year, "bt": bt, "avg": avg_top,
            "peak": peak, "s": n})
    return pid


def update_player(pid, **fields):
    allowed = {"name", "number", "position", "year", "bt", "avg_top_velo", "peak_velo", "active",
               "target_type", "target_date", "planned_pitches", "planned_innings", "sort"}
    for k, v in fields.items():
        if k not in allowed:
            continue
        if isinstance(v, date):
            v = v.isoformat()
        if isinstance(v, float) and pd.isna(v):
            v = None
        if k in ("avg_top_velo", "peak_velo", "planned_pitches", "planned_innings") and v is not None:
            v = float(v)
        if k in ("number",) and v is not None:
            v = str(v)
        _x(f"UPDATE players SET {k}=:v WHERE id=:id", {"v": v, "id": pid})


def remove_player(pid):
    """Hides his tab. His throw log is kept and comes back if he's added again."""
    update_player(pid, active=0)


# ───────── throw log (row level) ─────────
def _all_logs() -> pd.DataFrame:
    return _cached("logs", lambda: _df(f"SELECT id, player, created, {', '.join(LOG_FIELDS)} FROM throws"))


def load_log(pid) -> pd.DataFrame:
    df = _all_logs()
    df = df[df["player"] == pid].drop(columns=["player"])
    return df.sort_values(["date", "created"], ascending=False, na_position="last").reset_index(drop=True)


def get_row(row_id) -> dict | None:
    df = _all_logs()
    hit = df[df["id"] == row_id]
    if hit.empty:
        return None
    r = hit.iloc[0].to_dict()
    return {k: (None if (isinstance(v, float) and pd.isna(v)) else v) for k, v in r.items()}


def add_row(pid, rec: dict) -> str:
    rid = uuid.uuid4().hex[:12]
    now = time.time()
    params = {"id": rid, "player": pid, "created": rec.get("created", now), "updated": now}
    params.update({f: _clean(f, rec.get(f)) for f in LOG_FIELDS})
    _x(f"INSERT INTO throws(id, player, created, updated, {', '.join(LOG_FIELDS)}) "
       f"VALUES (:id, :player, :created, :updated, {', '.join(':' + f for f in LOG_FIELDS)})", params)
    return rid


def _add_rows_bulk(pid, recs: list[dict]):
    now = time.time()
    rows = []
    for i, rec in enumerate(recs):
        p = {"id": uuid.uuid4().hex[:12], "player": pid, "created": rec.get("created", 1000 + i), "updated": now}
        p.update({f: _clean(f, rec.get(f)) for f in LOG_FIELDS})
        rows.append(p)
    if rows:
        _x(f"INSERT INTO throws(id, player, created, updated, {', '.join(LOG_FIELDS)}) "
           f"VALUES (:id, :player, :created, :updated, {', '.join(':' + f for f in LOG_FIELDS)})", rows)


def update_row(row_id, rec: dict):
    fields = [f for f in LOG_FIELDS if f in rec]
    if not fields:
        return
    params = {f: _clean(f, rec[f]) for f in fields}
    params.update({"updated": time.time(), "id": row_id})
    _x(f"UPDATE throws SET {', '.join(f + '=:' + f for f in fields)}, updated=:updated WHERE id=:id", params)


def delete_row(row_id):
    _x("DELETE FROM throws WHERE id=:id", {"id": row_id})


# ───────── planner overrides ─────────
def load_overrides(pid) -> dict:
    df = _cached("overrides", lambda: _df("SELECT player, day, day_type FROM overrides"))
    df = df[df["player"] == pid]
    return {date.fromisoformat(d): t for d, t in zip(df["day"], df["day_type"])}


def save_override(pid, day: date, day_type):
    _x("DELETE FROM overrides WHERE player=:p AND day=:d", {"p": pid, "d": day.isoformat()})
    if day_type and day_type != "(none)":
        _x("INSERT INTO overrides VALUES (:p, :d, :t)", {"p": pid, "d": day.isoformat(), "t": day_type})


def clear_overrides(pid):
    _x("DELETE FROM overrides WHERE player=:p", {"p": pid})


# ───────── Build-Up plans ─────────
def load_buildup(pid) -> dict | None:
    df = _cached("buildup", lambda: _df("SELECT * FROM buildup"))
    hit = df[df["player"] == pid]
    if hit.empty:
        return None
    r = hit.iloc[0].to_dict()
    out = {}
    for k in BU_FIELDS:
        v = r.get(k)
        if isinstance(v, float) and pd.isna(v):
            v = None
        if k in ("start", "first_pen") and v:
            v = date.fromisoformat(v)
        out[k] = v
    return out


def all_buildups() -> dict:
    df = _cached("buildup", lambda: _df("SELECT * FROM buildup"))
    return {p: load_buildup(p) for p in df["player"]}


def save_buildup(pid, **fields):
    cur = load_buildup(pid) or {}
    cur.update({k: v for k, v in fields.items() if k in BU_FIELDS})
    params = {k: _clean(k, cur.get(k)) for k in BU_FIELDS}
    params["player"] = pid
    _x("DELETE FROM buildup WHERE player=:player", {"player": pid})
    _x(f"INSERT INTO buildup(player, {', '.join(BU_FIELDS)}) VALUES (:player, {', '.join(':' + k for k in BU_FIELDS)})", params)


def delete_buildup(pid):
    _x("DELETE FROM buildup WHERE player=:p", {"p": pid})
    _x("DELETE FROM buildup_days WHERE player=:p", {"p": pid})


def load_buildup_days(pid) -> dict:
    df = _cached("buildup_days", lambda: _df("SELECT * FROM buildup_days"))
    df = df[df["player"] == pid]
    out = {}
    for r in df.to_dict("records"):
        out[date.fromisoformat(r["day"])] = dict(
            off=bool(r["off"]), **{k: (None if pd.isna(r[k]) else float(r[k])) for k in ("throws", "distance", "velo_low", "velo_high")})
    return out


def save_buildup_day(pid, day: date, **fields):
    cur = load_buildup_days(pid).get(day, dict(off=False, throws=None, distance=None, velo_low=None, velo_high=None))
    cur.update(fields)
    _x("DELETE FROM buildup_days WHERE player=:p AND day=:d", {"p": pid, "d": day.isoformat()})
    if cur.get("off") or any(cur.get(k) not in (None, "") for k in ("throws", "distance", "velo_low", "velo_high")):
        _x("INSERT INTO buildup_days VALUES (:p, :d, :off, :t, :dist, :vl, :vh)",
           {"p": pid, "d": day.isoformat(), "off": 1 if cur.get("off") else 0, "t": _clean("throws", cur.get("throws")),
            "dist": _clean("distance", cur.get("distance")), "vl": _clean("velo_low", cur.get("velo_low")),
            "vh": _clean("velo_high", cur.get("velo_high"))})


def clear_buildup_days(pid):
    _x("DELETE FROM buildup_days WHERE player=:p", {"p": pid})


# ───────── first run / import ─────────
def is_empty() -> bool:
    return not _q("SELECT COUNT(*) FROM players")[0][0]


def seed_if_empty(workbook_path=None, backup_path=None):
    """First run: the live-app backup (roster, velo, every logged session) + the workbook's Build-Up plans."""
    if not is_empty():
        return False
    with _LOCK:
        if not is_empty():
            return False
        if backup_path and os.path.exists(backup_path):
            import_backup(backup_path)
        else:
            for num, name, pos, yr, bt in DEFAULT_ROSTER:
                add_player(name, num, pos, yr, bt)
            if workbook_path and os.path.exists(workbook_path):
                import_workbook(workbook_path, include_log=True)
        if workbook_path and os.path.exists(workbook_path):
            import_workbook(workbook_path, include_log=False, include_roster=False)
    return True


def import_backup(src) -> tuple[int, int]:
    """Re-load an Excel backup made by this app (Sidebar ▸ Import / backup ▸ Prepare Excel backup)."""
    x = pd.read_excel(src, sheet_name=None)
    ros = x.get("Roster")
    n_p = n_r = 0
    if ros is not None:
        for i, r in enumerate(ros.to_dict("records")):
            if pd.isna(r.get("name")):
                continue
            pid = r.get("id") if isinstance(r.get("id"), str) and r.get("id") else slugify(r["name"])
            num = r.get("number")
            num = "" if num is None or (isinstance(num, float) and pd.isna(num)) else str(int(num) if isinstance(num, float) else num)
            if not _q("SELECT 1 FROM players WHERE id=:id", {"id": pid}):
                _x("INSERT INTO players(id, name, active, sort, target_type) VALUES (:id, :n, 1, :s, 'Game')",
                   {"id": pid, "n": r["name"], "s": i + 1})
            fields = {k: r.get(k) for k in ("name", "position", "year", "bt", "avg_top_velo", "peak_velo",
                                            "target_type", "target_date", "planned_pitches", "planned_innings")}
            fields = {k: (None if (v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT) else v) for k, v in fields.items()}
            if isinstance(fields.get("target_date"), (pd.Timestamp, datetime)):
                fields["target_date"] = fields["target_date"].date()
            update_player(pid, number=num, sort=i + 1, active=1, **fields)
            n_p += 1
    log = x.get("Throw Log")
    if log is not None and len(log):
        log = log.reset_index(drop=True)
        by_player: dict = {}
        for i, r in enumerate(log.to_dict("records")):
            pid = resolve_player(r.get("player"))
            if not pid:
                pid = add_player(str(r.get("player")))
            rec = {f: r.get(f) for f in LOG_FIELDS}
            rec["created"] = 1000 + i
            by_player.setdefault(pid, []).append(rec)
        for pid, recs in by_player.items():
            _x("DELETE FROM throws WHERE player=:p", {"p": pid})   # replace, never duplicate
            recs.reverse()   # the backup lists newest first; keep the original entry order
            for j, rec in enumerate(recs):
                rec["created"] = 1000 + j
            _add_rows_bulk(pid, recs)
            n_r += len(recs)
    bset = x.get("Build-Up Settings")
    if bset is not None:
        for r in bset.to_dict("records"):
            pid = resolve_player(r.get("player"))
            if pid:
                save_buildup(pid, **{k: (None if pd.isna(r.get(k)) else r.get(k)) for k in BU_FIELDS})
    bdays = x.get("Build-Up Overrides")
    if bdays is not None:
        for r in bdays.to_dict("records"):
            pid = resolve_player(r.get("player"))
            if pid and not pd.isna(r.get("day")):
                save_buildup_day(pid, pd.Timestamp(r["day"]).date(), off=bool(r.get("off")) if not pd.isna(r.get("off")) else False,
                                 **{k: (None if pd.isna(r.get(k)) else r.get(k)) for k in ("throws", "distance", "velo_low", "velo_high")})
    return n_p, n_r


def import_workbook(src, include_log=True, include_roster=True) -> tuple[int, int, int]:
    """Import from the ACWR workbook: Roster velo, Throw Log, planner inputs and every Build-Up tab."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(src)
    n_players = n = n_bu = 0
    if include_roster:
        for row in wb["Roster"].iter_rows(min_row=5, values_only=True):
            if row[0]:
                name = str(row[0]).strip()
                pid = resolve_player(name) or add_player(name, position=row[1] or "RHP")
                fields = {"avg_top_velo": row[3]}
                if row[2] is not None:
                    fields["peak_velo"] = row[2]
                update_player(pid, **fields)
                n_players += 1
    if include_log:
        colmap = {1: "date", 3: "session_type", 4: "distance", 5: "start_time", 6: "end_time", 8: "actual_throws",
                  10: "velo_low", 11: "velo_high", 18: "warm_throws", 19: "warm_velo_low", 20: "warm_velo_high",
                  24: "sixty_throws", 25: "sixty_velo_low", 26: "sixty_velo_high", 30: "notes"}
        ws = wb["Throw Log"]
        by_player: dict = {}
        for r in range(5, ws.max_row + 1):
            name, stype = ws.cell(r, 2).value, ws.cell(r, 3).value
            if not name or not stype:      # skip blank placeholder rows
                continue
            rec = {}
            for col, k in colmap.items():
                v = ws.cell(r, col).value
                rec[k] = None if isinstance(v, str) and v.startswith("=") else v
            rec["created"] = 1000 + r
            pid = resolve_player(name) or add_player(str(name))
            by_player.setdefault(pid, []).append(rec)
        for pid, recs in by_player.items():
            _x("DELETE FROM throws WHERE player=:p", {"p": pid})
            _add_rows_bulk(pid, recs)
            n += len(recs)
        if "Next Outing Planner" in wb.sheetnames:
            nop = wb["Next Outing Planner"]
            pid = resolve_player(nop["B4"].value)
            if pid and isinstance(nop["B6"].value, datetime):
                update_player(pid, target_type=nop["B5"].value or "Game", target_date=nop["B6"].value.date(),
                              planned_pitches=nop["B23"].value, planned_innings=nop["D23"].value)
    for ws in wb.worksheets:
        if not ws.title.startswith("Build-Up"):
            continue
        pid = resolve_player(ws["D4"].value)
        if not pid or not isinstance(ws["D6"].value, datetime) or not isinstance(ws["D7"].value, datetime):
            continue
        start = ws["D6"].value.date()
        save_buildup(pid, start=start, first_pen=ws["D7"].value.date(), pen_pitches=ws["D8"].value,
                     pen_velo_low=ws["D9"].value, pen_velo_high=ws["E9"].value, starting_point=ws["D10"].value,
                     max_distance=ws["D11"].value, volume=ws["D12"].value)
        clear_buildup_days(pid)
        from datetime import timedelta
        for r in range(43, 184):
            j, k = ws[f"J{r}"].value, ws[f"K{r}"].value
            ao, ap, aq = ws[f"AO{r}"].value, ws[f"AP{r}"].value, ws[f"AQ{r}"].value
            if j == "x" or any(v not in (None, "") for v in (k, ao, ap, aq)):
                save_buildup_day(pid, start + timedelta(days=r - 43), off=(j == "x"), throws=k, distance=ao,
                                 velo_low=ap, velo_high=aq)
        n_bu += 1
    return n_players, n, n_bu


def export_backup(players_df, calc_frames: list[pd.DataFrame]) -> bytes:
    """Everything in one .xlsx — roster, throw log, Build-Up settings and day overrides. Re-importable."""
    import io
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        players_df.drop(columns=["active", "sort"], errors="ignore").to_excel(xw, sheet_name="Roster", index=False)
        (pd.concat(calc_frames) if calc_frames else pd.DataFrame()).to_excel(xw, sheet_name="Throw Log", index=False)
        names = dict(zip(players_df["id"], players_df["name"]))
        bu = [dict(player=names.get(p, p), **s) for p, s in all_buildups().items()]
        pd.DataFrame(bu, columns=["player"] + BU_FIELDS).to_excel(xw, sheet_name="Build-Up Settings", index=False)
        days = []
        for p in all_buildups():
            for d, o in load_buildup_days(p).items():
                days.append(dict(player=names.get(p, p), day=d, **o))
        pd.DataFrame(days, columns=["player", "day", "off", "throws", "distance", "velo_low", "velo_high"]).to_excel(
            xw, sheet_name="Build-Up Overrides", index=False)
    return buf.getvalue()
